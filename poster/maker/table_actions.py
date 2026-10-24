"""Bounded JSON tables, CSV and XLSX. Independent of SQLite and desktop code."""

import csv
from functools import partial
import io
import json
from typing import Any, Literal
import zipfile
from pydantic import Field, model_validator
from .actions import Action
from .action_fields import properties
from .data_actions import bounded, key, load_json
from .model import DataModel, GraphError
from .resources import MAX_DATA


class Table(DataModel):
    columns: list[str] = Field(default_factory=list, max_length=256)
    rows: list[list[Any]] = Field(default_factory=list, max_length=10000)

    @model_validator(mode="after")
    def structure(self):
        if len(set(self.columns)) != len(self.columns) or any(not c or len(c) > 200 for c in self.columns):
            raise ValueError("TABLE_INVALID")
        if any(len(r) != len(self.columns) for r in self.rows) or len(self.rows) * len(self.columns) > 100000:
            raise ValueError("TABLE_INVALID")
        bounded(self.model_dump())
        return self


def table(value):
    try:
        return Table.model_validate(value).model_dump()
    except (ValueError, TypeError) as exc:
        raise GraphError("TABLE_INVALID") from exc


def empty():
    return {"columns": [], "rows": []}


class TableArg(DataModel):
    table: Any = Field(
        default_factory=empty, description="Таблица {columns:[...],rows:[[...]]} / {{переменная}}"
    )


class Row(TableArg):
    row: list[Any] = Field(default_factory=list, max_length=256)


class RowIndex(TableArg):
    row: int = Field(default=0, ge=0, le=10000, strict=True)


class Cell(RowIndex):
    column: str = Field(default="value", description="Название столбца")


class SetCell(Cell):
    value: Any = None


class Find(TableArg):
    column: str = "value"
    operator: Literal[
        "equals", "not_equals", "gt", "ge", "lt", "le", "contains", "starts_with", "is_empty"
    ] = "equals"
    value: Any = None


class FromJSON(DataModel):
    value: Any = Field(default_factory=list, description="Массив объектов либо готовая таблица")


def matching(a, b, op):
    try:
        if op == "equals":
            return key(a) == key(b)
        if op == "not_equals":
            return key(a) != key(b)
        if op == "is_empty":
            return a is None or a == ""
        if op == "contains":
            return b in a
        if op == "starts_with":
            return isinstance(a, str) and a.startswith(str(b))
        if op == "gt":
            return a > b
        if op == "ge":
            return a >= b
        if op == "lt":
            return a < b
        return a <= b
    except TypeError:
        return False


async def table_action(ctx, p, operation):
    if operation == "Create":
        result = table(p.model_dump())
    elif operation == "FromJSON":
        if isinstance(p.value, dict):
            result = table(p.value)
        elif isinstance(p.value, list) and all(isinstance(r, dict) for r in p.value):
            columns = list(dict.fromkeys(k for row in p.value for k in row))
            result = table({"columns": columns, "rows": [[r.get(c) for c in columns] for r in p.value]})
        else:
            raise GraphError("TABLE_INVALID")
    else:
        value = table(p.table)
        result = value
        if operation == "ToJSON":
            result = [dict(zip(value["columns"], row, strict=True)) for row in value["rows"]]
        elif operation == "AddRow":
            value["rows"].append(p.row)
            result = table(value)
        elif operation == "RowCount":
            result = len(value["rows"])
        elif operation in {"GetCell", "SetCell", "RemoveRow"}:
            if p.row >= len(value["rows"]):
                raise GraphError("TABLE_CELL")
            if operation == "RemoveRow":
                value["rows"].pop(p.row)
            else:
                if p.column not in value["columns"]:
                    raise GraphError("TABLE_CELL")
                col = value["columns"].index(p.column)
                if operation == "GetCell":
                    result = value["rows"][p.row][col]
                else:
                    value["rows"][p.row][col] = p.value
                    result = table(value)
        else:
            if p.column not in value["columns"]:
                raise GraphError("TABLE_CELL")
            col = value["columns"].index(p.column)
            indices = [i for i, row in enumerate(value["rows"]) if matching(row[col], p.value, p.operator)]
            if operation == "FindRows":
                result = indices
            else:
                value["rows"] = [value["rows"][i] for i in indices]
    return "success", {"value": bounded(result)}


class CSVRead(DataModel):
    path: str = "input.csv"
    encoding: str = "utf-8-sig"
    delimiter: Literal[",", ";", "\t"] = ","
    has_header: bool = True
    infer_json: bool = False


class CSVWrite(TableArg):
    path: str = "output.csv"
    encoding: str = "utf-8-sig"
    delimiter: Literal[",", ";", "\t"] = ","
    header: bool = True
    overwrite: bool = False


def rows_table(rows, has_header):
    if not rows:
        return empty()
    if has_header and len(rows) == 1 and all(v is None for v in rows[0]):
        return empty()
    columns = (
        [str(v) if v is not None else "" for v in rows.pop(0)]
        if has_header
        else ["column_" + str(i + 1) for i in range(len(rows[0]))]
    )
    return table({"columns": columns, "rows": rows})


async def csv_action(ctx, p, operation):
    try:
        if operation == "Read":
            csv.field_size_limit(MAX_DATA)
            reader = csv.reader(
                io.StringIO(ctx.files.read(p.path).decode(p.encoding), newline=""),
                delimiter=p.delimiter,
                strict=True,
            )
            rows = []
            for row in reader:
                if len(rows) >= 10001 or len(row) > 256:
                    raise GraphError("DATA_LIMIT")
                if row:
                    rows.append(row)
            value = rows_table(rows, p.has_header)
            if p.infer_json:
                for row in value["rows"]:
                    for i, cell in enumerate(row):
                        try:
                            row[i] = load_json(cell)
                        except (ValueError, GraphError):
                            pass
        else:
            t = table(p.table)
            stream = io.StringIO(newline="")
            writer = csv.writer(stream, delimiter=p.delimiter, lineterminator="\r\n")
            if p.header:
                writer.writerow(t["columns"])
            for row in t["rows"]:
                writer.writerow(
                    [
                        ""
                        if v is None
                        else v
                        if isinstance(v, str)
                        else json.dumps(v, ensure_ascii=False, allow_nan=False)
                        for v in row
                    ]
                )
            value = ctx.files.write(p.path, stream.getvalue().encode(p.encoding), overwrite=p.overwrite)
        return "success", {"value": value}
    except (csv.Error, UnicodeError, LookupError, ValueError) as exc:
        if isinstance(exc, GraphError):
            raise
        raise GraphError("CSV_INVALID") from exc


class ExcelPath(DataModel):
    path: str = "data.xlsx"


class ExcelRead(ExcelPath):
    sheet: str = Field(
        default="", description="Пусто = первый лист; формулы не вычисляются, читается cached value"
    )
    has_header: bool = True


class ExcelWrite(TableArg):
    path: str = "data.xlsx"
    sheet: str = "Sheet1"
    update_existing: bool = Field(
        default=False, description="Явно разрешить изменение существующей простой XLSX-книги"
    )
    replace_sheet: bool = Field(default=False, description="Явно разрешить замену существующего листа")


def workbook_bytes(ctx, path, updating=False):
    if not path.lower().endswith(".xlsx"):
        raise GraphError("XLSX_INVALID")
    raw = ctx.files.read(path)
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 20_000_000:
                raise GraphError("XLSX_INVALID")
            for e in entries:
                if ".." in e.filename.split("/") or e.filename.startswith("/") or e.flag_bits & 1:
                    raise GraphError("XLSX_INVALID")
                if updating and (
                    e.filename.startswith(
                        (
                            "xl/drawings",
                            "xl/charts",
                            "xl/externalLinks",
                            "xl/pivot",
                            "xl/vba",
                            "xl/embeddings",
                            "customXml",
                            "xl/activeX",
                        )
                    )
                    or (e.filename.endswith(".xml") and b"extLst" in archive.read(e))
                ):
                    raise GraphError("XLSX_INVALID")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise GraphError("XLSX_INVALID") from exc
    return raw


async def excel_action(ctx, p, operation):
    import openpyxl
    from defusedxml.common import DefusedXmlException
    from xml.etree.ElementTree import ParseError

    book = None
    try:
        if operation in {"Read", "ListSheets"}:
            book = openpyxl.load_workbook(
                io.BytesIO(workbook_bytes(ctx, p.path)), read_only=True, data_only=True, keep_links=False
            )
            if operation == "ListSheets":
                value = book.sheetnames
            else:
                name = p.sheet or book.sheetnames[0]
                if name not in book.sheetnames:
                    raise GraphError("XLSX_INVALID")
                sheet = book[name]
                if (sheet.max_row or 0) > 10001 or (sheet.max_column or 0) > 256:
                    raise GraphError("DATA_LIMIT")
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if len(rows) >= 10001 or len(row) > 256:
                        raise GraphError("DATA_LIMIT")
                    # Excel dates are explicit ISO strings; JSON has no datetime type.
                    rows.append([v.isoformat() if hasattr(v, "isoformat") else v for v in row])
                value = rows_table(rows, p.has_header)
        else:
            t = table(p.table)
            path = ctx.files.path(p.path)
            if not p.path.lower().endswith(".xlsx"):
                raise GraphError("XLSX_INVALID")
            if path.exists():
                if not p.update_existing:
                    raise GraphError("FILE_EXISTS")
                book = openpyxl.load_workbook(io.BytesIO(workbook_bytes(ctx, p.path, True)), keep_links=False)
                if p.sheet in book.sheetnames:
                    if not p.replace_sheet:
                        raise GraphError("FILE_EXISTS")
                    index = book.sheetnames.index(p.sheet)
                    book.remove(book[p.sheet])
                    sheet = book.create_sheet(p.sheet, index)
                else:
                    sheet = book.create_sheet(p.sheet)
            else:
                book = openpyxl.Workbook()
                sheet = book.active
                sheet.title = p.sheet
            sheet.append(t["columns"])
            for row in t["rows"]:
                sheet.append(
                    [json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for v in row]
                )
            # Literal strings must never become executable spreadsheet formulas.
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            stream = io.BytesIO()
            book.save(stream)
            value = ctx.files.write(p.path, stream.getvalue(), overwrite=p.update_existing)
        return "success", {"value": value}
    except (ValueError, KeyError, TypeError, zipfile.BadZipFile, DefusedXmlException, ParseError) as exc:
        if isinstance(exc, GraphError):
            raise
        raise GraphError("XLSX_INVALID") from exc
    finally:
        if book is not None:
            book.close()


def register_table_actions(registry):
    for name, label, model in [
        ("Create", "Создать", Table),
        ("AddRow", "Добавить строку", Row),
        ("GetCell", "Прочитать ячейку", Cell),
        ("SetCell", "Изменить ячейку", SetCell),
        ("RemoveRow", "Удалить строку", RowIndex),
        ("FindRows", "Найти строки", Find),
        ("Filter", "Фильтр", Find),
        ("RowCount", "Количество строк", TableArg),
        ("FromJSON", "Из JSON", FromJSON),
        ("ToJSON", "В JSON", TableArg),
    ]:
        registry.register(
            Action(
                "Tables." + name,
                "Таблицы · " + label,
                model,
                partial(table_action, operation=name),
                properties(model),
                reveal_parameters=True,
            )
        )
    for family, name, label, model, execute in [
        ("CSV", "Read", "Прочитать CSV", CSVRead, csv_action),
        ("CSV", "Write", "Записать CSV", CSVWrite, csv_action),
        ("Excel", "Read", "Прочитать XLSX", ExcelRead, excel_action),
        ("Excel", "Write", "Записать XLSX", ExcelWrite, excel_action),
        ("Excel", "ListSheets", "Листы XLSX", ExcelPath, excel_action),
    ]:
        registry.register(
            Action(
                family + "." + name,
                "Таблицы · " + label,
                model,
                partial(execute, operation=name),
                properties(model, paths=("path",)),
                reveal_parameters=True,
            )
        )
